"""
tests/test_presisi_baru.py
==========================
Unit test untuk 3 perbaikan presisi baru:

1. pStyle Heading2 -> sub_bab, Heading3 -> sub_sub_bab
   (tanpa bergantung font_size yang bisa None karena inherited dari Word style)

2. pengesahan_nama: nama penanda tangan di-center via refinement kontekstual
   (paragraf 'isi' bertetangga pengesahan_tanda_tangan -> pengesahan_nama)

3. Formatter menghapus paragraf kosong manual sebelum sub_sub_bab (Heading3)
"""

import os
import tempfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from unittest.mock import MagicMock, patch

from services.rule_based_extractor import _heuristic_label
from services.structure_classifier import _refine_pengesahan_nama


# ============================================================
# PERBAIKAN 1: pStyle Heading2/3 -> sub_bab / sub_sub_bab
# ============================================================

@pytest.mark.parametrize("pstyle,expected", [
    ("Heading2",  "sub_bab"),
    ("Heading 2", "sub_bab"),
    ("heading 2", "sub_bab"),
    ("Heading3",  "sub_sub_bab"),
    ("Heading 3", "sub_sub_bab"),
    ("heading 3", "sub_sub_bab"),
])
def test_heading_pstyle_mapped_correctly(pstyle, expected):
    """
    Paragraf pStyle Heading2/3 harus di-label benar meskipun font_size=None
    (ukuran diwarisi dari Word style, bukan dari run).
    """
    para = MagicMock()
    mock_pstyle = MagicMock()
    mock_pstyle.val = pstyle
    para._p.pPr = MagicMock()
    para._p.pPr.pStyle = mock_pstyle

    with patch("services.rule_based_extractor.get_paragraph_text", return_value="Latar Belakang"), \
         patch("services.rule_based_extractor.get_effective_font_size", return_value=None), \
         patch("services.rule_based_extractor.get_effective_bold", return_value=None), \
         patch("services.rule_based_extractor.get_alignment_str", return_value="LEFT"), \
         patch("services.rule_based_extractor.has_numbering", return_value=False):
        label = _heuristic_label(para, para_index=80)
        assert label == expected, (
            "pStyle={!r} harus menghasilkan {!r}, dapat {!r}".format(pstyle, expected, label)
        )


# ============================================================
# PERBAIKAN 2: pengesahan_nama refinement kontekstual
# ============================================================

def test_refine_pengesahan_nama_detects_adjacent_name():
    """
    Paragraf 'isi' tepat sebelum/sesudah pengesahan_tanda_tangan
    di area lembar pengesahan harus -> pengesahan_nama.
    """
    doc = Document()
    doc.add_paragraph("LEMBAR PENGESAHAN")           # idx 0 -> pengesahan_heading
    doc.add_paragraph("")                              # idx 1 -> kosong
    doc.add_paragraph("Dodik Arwin Dermawan, S.ST.")  # idx 2 -> isi -> pengesahan_nama
    doc.add_paragraph("NIP. 197801082000121001")       # idx 3 -> pengesahan_tanda_tangan
    doc.add_paragraph("")                              # idx 4 -> kosong
    doc.add_paragraph("BAB I")                         # idx 5 -> judul_bab

    initial = {
        0: "pengesahan_heading",
        2: "isi",
        3: "pengesahan_tanda_tangan",
        5: "judul_bab",
    }
    refined = _refine_pengesahan_nama(doc, initial)
    assert refined[2] == "pengesahan_nama", (
        "Nama di sebelah NIP harus menjadi pengesahan_nama, dapat: " + repr(refined[2])
    )
    assert refined[3] == "pengesahan_tanda_tangan"
    assert refined[5] == "judul_bab"


def test_refine_pengesahan_nama_no_false_positive():
    """
    Tanpa pengesahan_heading, tidak ada reklasifikasi yang terjadi.
    """
    doc = Document()
    doc.add_paragraph("Ini paragraf isi biasa")
    doc.add_paragraph("NIP. 197801082000121001")
    doc.add_paragraph("BAB I PENDAHULUAN")

    initial = {
        0: "isi",
        1: "pengesahan_tanda_tangan",
        2: "judul_bab",
    }
    refined = _refine_pengesahan_nama(doc, initial)
    assert refined[0] == "isi"  # tidak berubah karena di luar area pengesahan


# ============================================================
# PERBAIKAN 3: Hapus baris kosong sebelum Heading3 (sub_sub_bab)
# ============================================================

def test_blank_para_before_heading3_removed():
    """
    Formatter harus menghapus paragraf kosong yang langsung mendahului
    paragraf pStyle=Heading3 sebelum proses format.
    """
    from services.rule_based_formatter import format_document

    doc = Document()
    doc.add_paragraph("Isi biasa")
    doc.add_paragraph("")      # blank yang harus dihapus
    h3 = doc.add_paragraph("Subbab Level 3")

    # Set pStyle ke Heading3
    p_pr = h3._p.get_or_add_pPr()
    pStyle_el = OxmlElement("w:pStyle")
    pStyle_el.set(qn("w:val"), "Heading3")
    p_pr.insert(0, pStyle_el)

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_in:
        doc.save(tmp_in.name)
        input_path = tmp_in.name

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_out:
        output_path = tmp_out.name

    try:
        cfg = {
            "margins": {},
            "styles": {"isi": {"line_spacing": 1.5}},
            "default_line_spacing": 1.5,
        }
        format_document(input_path, output_path, cfg)

        result_doc = Document(output_path)
        texts = [p.text.strip() for p in result_doc.paragraphs]

        assert "Subbab Level 3" in texts
        idx_isi = texts.index("Isi biasa") if "Isi biasa" in texts else -1
        idx_h3  = texts.index("Subbab Level 3")

        if idx_isi >= 0 and idx_h3 > idx_isi + 1:
            between = texts[idx_isi + 1:idx_h3]
            non_empty_between = [t for t in between if t]
            assert non_empty_between == [], (
                "Tidak boleh ada paragraf non-kosong di antara 'Isi biasa' dan 'Subbab Level 3': "
                + repr(between)
            )
            # Dan tidak boleh ada blank para sama sekali (blank sudah dihapus)
            assert between == [], (
                "Paragraf kosong sebelum Heading3 harus dihapus, tersisa: " + repr(between)
            )
    finally:
        os.unlink(input_path)
        try:
            os.unlink(output_path)
        except FileNotFoundError:
            pass
